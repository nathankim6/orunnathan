# -*- coding: utf-8 -*-
"""옳은영어 수능 유형 문제 세트 — items.py의 출제 데이터를 A4 워드 문서로 만든다."""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from split import load                      # noqa: E402
from items import ITEMS, TYPE_ORDER, TYPES  # noqa: E402
from balance import balance, ORDER_MAPS, ORDER_LABELS  # noqa: E402

# 팔레트 (옳은영어 v2)
INK    = RGBColor(0x1C, 0x1C, 0x1C)
BODY   = RGBColor(0x3A, 0x3A, 0x3A)
GREY   = RGBColor(0x8A, 0x8A, 0x8A)
BLUE   = RGBColor(0x1A, 0x7F, 0xBF)
GOLD   = RGBColor(0xD9, 0xB3, 0x00)
YELLOW = RGBColor(0xFF, 0xD4, 0x00)
HAIR   = 'E4E2DD'
PAPER  = 'F7F6F2'
LBLUE  = 'EAF4FA'
FONT   = 'Noto Sans KR'
CIRCLE = ['①', '②', '③', '④', '⑤']


# ---------------------------------------------------------------- 서식 헬퍼
# OOXML은 자식 요소 순서가 규격으로 정해져 있다. 순서를 어기면 워드·리브레오피스가
# 파일을 열지 못하므로, 항상 _tag_seq를 보고 제자리에 끼워 넣는다.
# python-docx는 클래스 정의가 끝나면 _tag_seq를 지우므로 직접 들고 있는다.
EDGE_SEQ = tuple('w:' + e for e in (
    'top', 'left', 'bottom', 'right', 'between', 'bar',
    'insideH', 'insideV', 'tl2br', 'tr2bl'))
RPR_SEQ = (
    'w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps', 'w:smallCaps',
    'w:strike', 'w:dstrike', 'w:outline', 'w:shadow', 'w:emboss', 'w:imprint',
    'w:noProof', 'w:snapToGrid', 'w:vanish', 'w:webHidden', 'w:color', 'w:spacing',
    'w:w', 'w:kern', 'w:position', 'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect',
    'w:bdr', 'w:shd', 'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang',
    'w:eastAsianLayout', 'w:specVanish', 'w:oMath')
PPR_SEQ = (
    'w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr',
    'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd', 'w:tabs',
    'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
    'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd',
    'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
    'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
    'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr',
    'w:sectPr', 'w:pPrChange')
TCPR_SEQ = (
    'w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge', 'w:tcBorders',
    'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
    'w:hideMark', 'w:headers', 'w:cellIns', 'w:cellDel', 'w:cellMerge', 'w:tcPrChange')
TBLPR_SEQ = (
    'w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual', 'w:tblStyleRowBandSize',
    'w:tblStyleColBandSize', 'w:tblW', 'w:jc', 'w:tblCellSpacing', 'w:tblInd',
    'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
    'w:tblCaption', 'w:tblDescription', 'w:tblPrChange')
SEQ_BY_PARENT = {'rPr': RPR_SEQ, 'pPr': PPR_SEQ, 'tcPr': TCPR_SEQ, 'tblPr': TBLPR_SEQ,
                 'tcBorders': EDGE_SEQ, 'pBdr': EDGE_SEQ, 'tblBorders': EDGE_SEQ}


def ordered(parent, tag, seq=None):
    """parent 안에서 tag를 찾고, 없으면 규격 순서에 맞춰 새로 넣는다."""
    found = parent.find(qn(tag))
    if found is not None:
        return found
    seq = seq or SEQ_BY_PARENT.get(parent.tag.rsplit('}', 1)[-1])
    el = OxmlElement(tag)
    if seq and tag in seq:
        later = {qn(t) for t in seq[seq.index(tag) + 1:]}
        for child in parent:
            if child.tag in later:
                child.addprevious(el)
                return el
    parent.append(el)
    return el


def _fonts(run, name=FONT):
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), name)


def run(p, text, size=10, bold=False, color=BODY, italic=False,
        underline=False, spacing=None, name=FONT):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.underline = underline
    r.font.color.rgb = color
    _fonts(r, name)
    if spacing:  # 자간
        ordered(r._element.get_or_add_rPr(), 'w:spacing').set(qn('w:val'), str(spacing))
    return r


def para(container, text='', size=10, bold=False, color=BODY, align=None,
         before=0, after=4, line=1.35, indent=0, keep=False, **kw):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    if text:
        run(p, text, size=size, bold=bold, color=color, **kw)
    return p


def shade(el, fill):
    sh = ordered(el, 'w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)


def cell_shade(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def borders(el, edges, color=HAIR, sz=4, val='single'):
    """edges: 'top bottom left right' 중 필요한 것만. el은 tcPr 또는 pPr."""
    tag = 'w:tcBorders' if el.tag.endswith('}tcPr') else 'w:pBdr'
    bd = ordered(el, tag)
    for e in edges.split():
        node = ordered(bd, 'w:' + e)
        node.set(qn('w:val'), val)
        node.set(qn('w:sz'), str(sz))
        node.set(qn('w:space'), '4')
        node.set(qn('w:color'), color)


def no_table_borders(table):
    bd = ordered(table._tbl.tblPr, 'w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        n = ordered(bd, 'w:' + e)
        n.set(qn('w:val'), 'none')
        n.set(qn('w:sz'), '0')
    return bd


def hairline(doc, before=2, after=6):
    p = para(doc, '', before=before, after=after, line=1.0)
    borders(p._p.get_or_add_pPr(), 'bottom')
    return p


def field(p, instr):
    """PAGE / NUMPAGES 필드. begin·instrText·end를 각각 별도 run으로 넣는다."""
    def _r():
        r = p.add_run()
        r.font.size = Pt(7.5)
        r.font.color.rgb = GREY
        _fonts(r)
        return r
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
    _r()._r.append(a)
    t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr
    _r()._r.append(t)
    z = OxmlElement('w:fldChar'); z.set(qn('w:fldCharType'), 'end')
    _r()._r.append(z)


# ---------------------------------------------------------------- 머리말/꼬리말
def build_header(section, subtitle):
    hp = section.header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(2)
    hp.paragraph_format.tab_stops.add_tab_stop(Emu(int(9638 * 635)), WD_ALIGN_PARAGRAPH.RIGHT)
    r = hp.add_run()
    r.add_picture(os.path.join(HERE, '..', 'assets', 'logo.png'), height=Pt(11))
    run(hp, '  ORUN ENGLISH', size=8.5, bold=True, color=INK, spacing=40)
    run(hp, '\t' + subtitle, size=8, color=GREY)
    borders(hp._p.get_or_add_pPr(), 'bottom')

    fp = section.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Emu(int(9638 * 635)), WD_ALIGN_PARAGRAPH.RIGHT)
    borders(fp._p.get_or_add_pPr(), 'top')
    run(fp, 'ORUN ENGLISH', size=7.5, bold=True, color=GREY, spacing=40)
    run(fp, '  옳은영어 · 고등 영어 수능 유형 문제 세트', size=7.5, color=GREY)
    run(fp, '\t', size=7.5, color=GREY)
    field(fp, 'PAGE')
    run(fp, ' / ', size=7.5, color=GREY)
    field(fp, 'NUMPAGES')


# ---------------------------------------------------------------- 지문 조판
def passage_para(doc, text, underline=None, blank=None, indent=0.35):
    """지문 한 덩어리. underline 구문은 밑줄, blank 구문은 빈칸으로 바꾼다."""
    p = para(doc, '', after=6, line=1.45, indent=indent, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    chunks = [(text, {})]
    if underline:
        assert underline in text, f'밑줄 구문 없음: {underline[:40]}'
        a, b = text.split(underline, 1)
        chunks = [(a, {}), (underline, {'underline': True, 'bold': True}), (b, {})]
    elif blank:
        assert blank in text, f'빈칸 구문 없음: {blank[:40]}'
        a, b = text.split(blank, 1)
        chunks = [(a, {}), ('_' * 34, {}), (b, {})]
    for t, kw in chunks:
        if t:
            run(p, t, size=9.5, **kw)
    return p


def boxed(doc, lines, fill=PAPER, size=9.5, label=None):
    """1셀 표 콜아웃(주어진 글·주어진 문장·요약문)."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    no_table_borders(t)
    c = t.cell(0, 0)
    c.width = Cm(16.4)
    cell_shade(c, fill)
    c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
    if label:
        p = c.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run(p, label, size=7.5, bold=True, color=GOLD if fill == PAPER else BLUE, spacing=20)
    for i, ln in enumerate(lines):
        p = c.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0 if i == 0 and not label else 3)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.4
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run(p, ln, size=size)
    para(doc, '', after=3, line=1.0)
    return t


def choices(doc, opts, korean=False, inline=False):
    if inline:  # 순서·문장삽입형: 한 줄에 여러 개
        p = para(doc, '', after=6, before=2, indent=0.35)
        for i, o in enumerate(opts):
            run(p, (CIRCLE[i] + ' ' + o).strip(), size=9.5)
            if i != len(opts) - 1:
                run(p, '        ' if not o else '     ', size=9.5)
        return
    for i, o in enumerate(opts):
        p = para(doc, '', after=1, before=0, line=1.3, indent=0.35)
        pf = p.paragraph_format
        pf.first_line_indent = Cm(-0.55)
        pf.left_indent = Cm(0.9)
        run(p, CIRCLE[i] + ' ', size=9.5, bold=True, color=INK)
        run(p, o, size=9.5)
    para(doc, '', after=2, line=1.0)


def summary_choices(doc, pairs):
    t = doc.add_table(rows=len(pairs) + 1, cols=3)
    t.autofit = False
    no_table_borders(t)
    widths = (Cm(1.0), Cm(6.2), Cm(6.2))
    hdr = ('', '(A)', '(B)')
    for j in range(3):
        c = t.cell(0, j)
        c.width = widths[j]
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        run(p, hdr[j], size=8.5, bold=True, color=GREY, spacing=20)
        borders(c._tc.get_or_add_tcPr(), 'bottom', color='1C1C1C', sz=8)
    for i, (a, b) in enumerate(pairs, start=1):
        for j, txt in enumerate((CIRCLE[i - 1], a, b)):
            c = t.cell(i, j)
            c.width = widths[j]
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            run(p, txt, size=9.5, bold=(j == 0), color=INK if j == 0 else BODY)
        borders(t.cell(i, 0)._tc.get_or_add_tcPr(), 'bottom')
        borders(t.cell(i, 1)._tc.get_or_add_tcPr(), 'bottom')
        borders(t.cell(i, 2)._tc.get_or_add_tcPr(), 'bottom')
    para(doc, '', after=3, line=1.0)


# ---------------------------------------------------------------- 문항 변환
def order_blocks(sents, lead, cuts, ans):
    """정답 번호에 맞춰 (A)(B)(C) 라벨을 붙인 블록을 돌려준다."""
    rest = sents[lead:]
    i, j = cuts
    c1, c2, c3 = rest[:i], rest[i:j], rest[j:]
    assert all((c1, c2, c3)), '순서 블록이 비었음'
    labels = ORDER_MAPS[ans]
    tagged = dict(zip(labels, (' '.join(c1), ' '.join(c2), ' '.join(c3))))
    return ' '.join(sents[:lead]), [(k, tagged[k]) for k in ('A', 'B', 'C')]


def insert_text(sents, take, ans):
    """take번째 문장을 빼고, 정답이 ans번이 되도록 ( ① )~( ⑤ )를 놓는다."""
    assert 1 <= take < len(sents), 'take 범위 오류'
    rest = sents[:take] + sents[take + 1:]
    start = take - ans
    assert start >= 0 and start + 4 <= len(rest) - 1, \
        f'삽입 위치 배치 불가 (문장 {len(sents)}, take {take}, 정답 {ans})'
    out = []
    for k, s in enumerate(rest):
        out.append(s)
        gap = k - start + 1
        if 1 <= gap <= 5:
            out.append(f'( {CIRCLE[gap - 1]} )')
    return sents[take], ' '.join(out)


# ---------------------------------------------------------------- 문항 출력
def render_question(doc, no, pid, sents, passage, it, kind):
    d = it[kind]
    label = TYPES[kind]['label']
    stem = TYPES[kind]['stem']
    if kind == 'implication':
        stem = f'밑줄 친 “{d["phrase"]}”이(가) 다음 글에서 의미하는 바로 가장 적절한 것은?'

    p = para(doc, '', before=10, after=5, keep=True)
    run(p, f'{no:03d}', size=11, bold=True, color=INK)
    run(p, f'  [{label}]  ', size=7.5, bold=True, color=GOLD, spacing=20)
    run(p, stem, size=9.5, bold=True, color=INK)

    if kind == 'order':
        lead, blocks = order_blocks(sents, d['lead'], d['cuts'], d['ans'])
        boxed(doc, [lead], label='주어진 글')
        for tag, txt in blocks:
            pp = para(doc, '', after=5, line=1.45, indent=0.35,
                      align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            run(pp, f'({tag}) ', size=9.5, bold=True, color=INK)
            run(pp, txt, size=9.5)
        choices(doc, ORDER_LABELS[:3], inline=True)
        p2 = para(doc, '', after=6, before=0, indent=0.35)
        for i in (3, 4):
            run(p2, f'{CIRCLE[i]} {ORDER_LABELS[i]}', size=9.5)
            if i == 3:
                run(p2, '     ', size=9.5)
        return
    if kind == 'insert':
        given, body = insert_text(sents, d['take'], d['ans'])
        boxed(doc, [given], fill=LBLUE, label='주어진 문장')
        passage_para(doc, body)
        choices(doc, [''] * 5, inline=True)
        return

    if kind == 'implication':
        passage_para(doc, passage, underline=d['phrase'])
    elif kind == 'blank':
        passage_para(doc, passage, blank=d['target'])
    else:
        passage_para(doc, passage)

    if kind == 'summary':
        boxed(doc, [d['tmpl']], label='요약문')
        summary_choices(doc, d['ch'])
    else:
        choices(doc, d['ch'], korean=(kind == 'mainPoint'))


def render_solution(doc, no, pid, sents, passage, it, kind):
    d = it[kind]
    p = para(doc, '', before=7, after=2, keep=True)
    run(p, f'{no:03d}', size=10, bold=True, color=INK)
    run(p, f'  {TYPES[kind]["label"]}', size=7.5, bold=True, color=GOLD, spacing=20)
    run(p, '   정답 ', size=9, color=GREY)
    ans = d['ans']
    run(p, CIRCLE[ans - 1], size=11, bold=True, color=BLUE)

    if kind == 'order':
        run(p, f'   {ORDER_LABELS[ans - 1]}', size=9, color=GREY)
    if kind == 'insert':
        run(p, f'   ({d["take"]+1}번째 문장 자리)', size=9, color=GREY)

    para(doc, d['sol'], size=9, after=3, line=1.4, indent=0.35,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    if kind in ('blank', 'implication'):
        key = 'target' if kind == 'blank' else 'phrase'
        q = para(doc, '', size=9, after=3, indent=0.35, line=1.4)
        run(q, '원문 표현  ', size=8, bold=True, color=GREY, spacing=20)
        run(q, d[key], size=9, italic=True)

    if d.get('tr'):
        para(doc, '보기 해석', size=8, bold=True, color=GREY, after=1,
             indent=0.35, spacing=20)
        for i, t in enumerate(d['tr']):
            pp = para(doc, '', after=0, line=1.3, indent=0.35)
            pp.paragraph_format.first_line_indent = Cm(-0.55)
            pp.paragraph_format.left_indent = Cm(0.9)
            run(pp, CIRCLE[i] + ' ', size=8.5, bold=True,
                color=BLUE if i == ans - 1 else GREY)
            run(pp, t, size=8.5, color=BODY if i == ans - 1 else GREY)
        para(doc, '', after=2, line=1.0)


# ---------------------------------------------------------------- 문서 조립
def h1(doc, num, text):
    p = para(doc, '', before=16, after=6, keep=True)
    run(p, '■ ', size=7, color=YELLOW)
    run(p, f'{num}   ', size=8.5, bold=True, color=GOLD, spacing=60)
    run(p, text, size=13.5, bold=True, color=INK)
    hairline(doc, before=0, after=4)


def info_table(doc, rows, widths=(Cm(3.4), Cm(13.0))):
    t = doc.add_table(rows=len(rows), cols=2)
    t.autofit = False
    no_table_borders(t)
    for i, (k, v) in enumerate(rows):
        for j, txt in enumerate((k, v)):
            c = t.cell(i, j)
            c.width = widths[j]
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.3
            if j == 0:
                run(p, '■ ', size=6.5, color=YELLOW if i % 2 == 0 else BLUE)
                run(p, txt, size=9, bold=True, color=GREY, spacing=20)
            else:
                run(p, txt, size=9.5)
            borders(c._tc.get_or_add_tcPr(), 'bottom')
    para(doc, '', after=4, line=1.0)
    return t


def build(out_path):
    rows = [r for r in load() if ITEMS.get(r['id'])]
    missing = [r['id'] for r in load() if not ITEMS.get(r['id'])]
    if missing:
        print('출제 데이터 없음:', ', '.join(missing))
    balance(ITEMS, rows)
    doc = Document()
    st = doc.styles['Normal']
    st.font.name, st.font.size = FONT, Pt(10)
    st.font.color.rgb = BODY
    rf = st.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), FONT)

    # python-docx 기본 템플릿의 <w:zoom>에 percent가 없어 스키마 검증에 걸린다.
    z = doc.settings.element.find(qn('w:zoom'))
    if z is not None and z.get(qn('w:percent')) is None:
        z.set(qn('w:percent'), '100')

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    build_header(sec, '옳은영어 · 수능 유형 문제 세트')

    # ── 표지
    para(doc, '', after=10, line=1.0)
    p = para(doc, '', after=4)
    run(p, '■ ', size=7, color=YELLOW)
    run(p, 'SUNEUNG TYPE PROBLEM SET', size=8, color=GREY, spacing=40)
    para(doc, '수능 유형 문제 세트', size=21, bold=True, color=INK, after=2, line=1.1)
    para(doc, '가짜 뉴스 교과서 지문 4편 + 최근 모의평가·수능 기출 지문 21편', size=10, color=GREY, after=12)

    total = sum(len([k for k in TYPE_ORDER if k in ITEMS[r['id']]]) for r in rows)
    info_table(doc, [
        ('지문', f'{len(rows)}편 (교과서 4편 · 2025학년도 9월 8편 · 수능 1편 · 2022학년도 9월 12편)'),
        ('유형', ' · '.join(TYPES[k]['label'] for k in TYPE_ORDER)),
        ('문항 수', f'총 {total}문항 (지문당 최대 {len(TYPE_ORDER)}문항)'),
        ('구성', '문제편 → 정답 및 해설편(지문 해석 포함)'),
    ])

    para(doc, '', after=6, line=1.0)
    p = para(doc, '', after=4)
    run(p, '■ ', size=7, color=YELLOW)
    run(p, '유형별 문항 수', size=8.5, bold=True, color=GOLD, spacing=40)
    tt = doc.add_table(rows=len(TYPE_ORDER) + 1, cols=3)
    tt.autofit = False
    no_table_borders(tt)
    W = (Cm(3.0), Cm(10.4), Cm(3.0))
    for j, h in enumerate(('수능 번호', '유형', '문항 수')):
        c = tt.cell(0, j); c.width = W[j]
        run(c.paragraphs[0], h, size=8.5, bold=True, color=GREY, spacing=20)
        borders(c._tc.get_or_add_tcPr(), 'bottom', color='1C1C1C', sz=8)
    for i, k in enumerate(TYPE_ORDER, start=1):
        n = sum(1 for r in rows if k in ITEMS[r['id']])
        for j, txt in enumerate((TYPES[k]['no'], TYPES[k]['name'], f'{n}문항')):
            c = tt.cell(i, j); c.width = W[j]
            pp = c.paragraphs[0]
            pp.paragraph_format.space_before = Pt(2)
            pp.paragraph_format.space_after = Pt(2)
            if j == 0:
                run(pp, '■ ', size=6.5, color=YELLOW if i % 2 else BLUE)
                run(pp, txt, size=9, bold=True)
            else:
                run(pp, txt, size=9.5)
            borders(c._tc.get_or_add_tcPr(), 'bottom')

    # ── 문제편
    doc.add_page_break()
    h1(doc, 'PART 1', '문제편')
    no = 0
    numbering = {}
    for idx, r in enumerate(rows, start=1):
        it = ITEMS[r['id']]
        kinds = [k for k in TYPE_ORDER if k in it]
        if idx > 1:
            doc.add_page_break()
        p = para(doc, '', before=0, after=2, keep=True)
        run(p, '■ ', size=7, color=YELLOW)
        run(p, f'지문 {idx:02d}', size=8.5, bold=True, color=GOLD, spacing=40)
        run(p, f'   {r["id"]}', size=11, bold=True, color=INK)
        if r.get('src_title'):
            run(p, f'   {r["src_title"]}', size=9, color=GREY)
        run(p, f'   ({len(kinds)}문항)', size=8.5, color=GREY)
        hairline(doc, before=0, after=6)
        for k in kinds:
            no += 1
            numbering[(r['id'], k)] = no
            render_question(doc, no, r['id'], r['sents'], r['en'], it, k)
        if len(kinds) < len(TYPE_ORDER):
            miss = [TYPES[k]['label'] for k in TYPE_ORDER if k not in it]
            para(doc, '※ ' + ', '.join(miss) + ' 유형은 이 지문의 문장 수가 모자라 출제하지 않았습니다.',
                 size=8, color=GREY, before=8, after=0)

    # ── 해설편
    doc.add_page_break()
    h1(doc, 'PART 2', '정답 및 해설')
    for idx, r in enumerate(rows, start=1):
        it = ITEMS[r['id']]
        kinds = [k for k in TYPE_ORDER if k in it]
        if idx > 1:
            doc.add_page_break()
        p = para(doc, '', before=0, after=2, keep=True)
        run(p, '■ ', size=7, color=YELLOW)
        run(p, f'지문 {idx:02d}', size=8.5, bold=True, color=GOLD, spacing=40)
        run(p, f'   {r["id"]}', size=11, bold=True, color=INK)
        rng = f'{numbering[(r["id"], kinds[0])]:03d}–{numbering[(r["id"], kinds[-1])]:03d}'
        run(p, f'   {rng}번', size=9, color=GREY)
        hairline(doc, before=0, after=5)
        boxed(doc, [r['ko']], fill=PAPER, size=9, label='지문 해석')
        for k in kinds:
            render_solution(doc, numbering[(r['id'], k)], r['id'], r['sents'], r['en'], it, k)

    doc.save(out_path)
    return no, out_path


if __name__ == '__main__':
    n, path = build(os.path.join(HERE, '..', '옳은영어_수능유형_문제세트.docx'))
    print(f'{n}문항 -> {path}')
